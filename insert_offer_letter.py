"""
Script to replace the offer letter text boxes on Slide 5 with the actual offer letter image.
"""
import shutil
import os
from pptx import Presentation
from pptx.util import Emu
from lxml import etree

# ── Config ──────────────────────────────────────────────────────────────────
PPTX_IN  = "Intership persentation.pptx"
PPTX_OUT = "Intership persentation.pptx"
IMG_PATH = "offer_letter.jpg"   # The saved offer letter photo
SLIDE_INDEX = 4                 # Slide 5 (0-based)

# Shapes to keep on Slide 5 (these are the slide-template / header shapes)
KEEP_SHAPE_NAMES = {
    "TextBox 1",   # AMITY UNIVERSITY
    "TextBox 2",   # PATNA
    "Rectangle 3", # logo box
    "TextBox 4",   # "A"
    "Rectangle 5", # top bar left
    "Rectangle 6", # top bar right
    "Rectangle 7", # bottom bar
    "TextBox 8",   # "Offer Letter" heading
}

# The bounding box of the offer letter content area (Rectangle 9 coords)
# left=2148840, top=1627632, width=6675120, height=5074920
CONTENT_LEFT   = 2148840
CONTENT_TOP    = 1627632
CONTENT_WIDTH  = 6675120
CONTENT_HEIGHT = 5074920

# ── Backup ──────────────────────────────────────────────────────────────────
backup = PPTX_IN.replace(".pptx", "_backup.pptx")
if not os.path.exists(backup):
    shutil.copy2(PPTX_IN, backup)
    print(f"Backup saved: {backup}")
else:
    print(f"Backup already exists: {backup}")

# ── Open presentation ────────────────────────────────────────────────────────
prs  = Presentation(PPTX_IN)
slide = prs.slides[SLIDE_INDEX]
sp_tree = slide.shapes._spTree   # XML element containing all shapes

# ── Remove unwanted shapes ───────────────────────────────────────────────────
to_remove = []
for shape in slide.shapes:
    if shape.name not in KEEP_SHAPE_NAMES:
        to_remove.append(shape._element)
        print(f"  Removing: {shape.name}  ({getattr(shape, 'text', '')}...)"[:80])

for elem in to_remove:
    sp_tree.remove(elem)

print(f"\nRemoved {len(to_remove)} shapes.")

# ── Insert image ─────────────────────────────────────────────────────────────
pic = slide.shapes.add_picture(
    IMG_PATH,
    left   = CONTENT_LEFT,
    top    = CONTENT_TOP,
    width  = CONTENT_WIDTH,
    height = CONTENT_HEIGHT,
)
print(f"Inserted image: {IMG_PATH}")

# ── Save ─────────────────────────────────────────────────────────────────────
prs.save(PPTX_OUT)
print(f"\nSaved: {PPTX_OUT}")
