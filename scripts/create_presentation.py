from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "quickCart_Project_Presentation_Ayush.pptx"
LOGO = ROOT / "assets" / "images" / "quickcart-logo.png"
BANNER = ROOT / "assets" / "images" / "Screenshot_2026-07-30_000109.png"

NAVY = RGBColor(0, 45, 92)
GOLD = RGBColor(255, 204, 0)
BLACK = RGBColor(0, 0, 0)
GREEN = RGBColor(34, 197, 94)
GRAY = RGBColor(245, 247, 250)
DARK = RGBColor(25, 25, 25)


def add_text(slide, text, x, y, w, h, size=24, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=21):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.space_after = Pt(10)
        p.bullet = True
    return box


def add_header(slide, title=None):
    add_text(slide, "AMITY\nUNIVERSITY", 1.18, 0.22, 1.45, 0.72, 23, True, NAVY, PP_ALIGN.CENTER)
    add_text(slide, "PATNA", 1.42, 0.85, 0.9, 0.18, 9, False, NAVY, PP_ALIGN.CENTER)
    shield = slide.shapes.add_shape(1, Inches(0.72), Inches(0.27), Inches(0.42), Inches(0.58))
    shield.fill.solid()
    shield.fill.fore_color.rgb = GOLD
    shield.line.color.rgb = NAVY
    add_text(slide, "A", 0.78, 0.39, 0.3, 0.2, 14, True, NAVY, PP_ALIGN.CENTER)
    slide.shapes.add_shape(1, Inches(0.72), Inches(1.08), Inches(5.6), Inches(0.06)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = GOLD
    slide.shapes.add_shape(1, Inches(6.32), Inches(1.08), Inches(4.95), Inches(0.06)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLACK
    slide.shapes.add_shape(1, Inches(0), Inches(7.42), Inches(12), Inches(0.08)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = NAVY
    if title:
        add_text(slide, title, 4.0, 1.36, 4.0, 0.5, 26, True, BLACK, PP_ALIGN.CENTER)


def add_title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s)
    add_text(s, "quickCart", 4.2, 1.8, 3.6, 0.55, 30, True, BLACK, PP_ALIGN.CENTER)
    add_text(s, "A Hyperlocal Quick-Commerce Mobile Application", 3.1, 2.35, 5.8, 0.38, 18, False, DARK, PP_ALIGN.CENTER)
    add_text(s, "Presented By", 1.45, 3.0, 3.0, 0.35, 18, True, BLACK, PP_ALIGN.CENTER)
    add_text(s, "Ayush (A45304825002) - BCA", 1.05, 3.45, 3.8, 0.35, 18, False, BLACK, PP_ALIGN.CENTER)
    add_text(s, "Under the Supervision of", 7.25, 3.0, 3.2, 0.35, 18, True, BLACK, PP_ALIGN.CENTER)
    add_text(s, "Dr. Sushant Kumar Dubey", 7.1, 3.45, 3.55, 0.35, 18, False, BLACK, PP_ALIGN.CENTER)
    add_text(s, "Amity Institute Of Information Technology, Patna\n10 August, 2026", 3.1, 5.55, 5.8, 0.6, 18, False, BLACK, PP_ALIGN.CENTER)


def simple_slide(prs, title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, title)
    add_bullets(s, bullets, 1.15, 2.25, 9.8, 4.3)
    return s


def add_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, "System Architecture")
    labels = [("Mobile App\nExpo + React Native", 0.9, 2.55), ("Services Layer\nOrders, Address,\nNotifications", 4.0, 2.55), ("Firebase\nAuth, Firestore,\nStorage", 7.1, 2.55), ("Admin Panel\nInventory, Users,\nOrders", 4.0, 4.75)]
    for label, x, y in labels:
        shape = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.35), Inches(0.9))
        shape.fill.solid()
        shape.fill.fore_color.rgb = GRAY
        shape.line.color.rgb = NAVY
        add_text(s, label, x + 0.08, y + 0.12, 2.18, 0.62, 15, True, DARK, PP_ALIGN.CENTER)
    for x1, y1, x2, y2 in [(3.25, 3.0, 4.0, 3.0), (6.35, 3.0, 7.1, 3.0), (5.18, 3.45, 5.18, 4.75)]:
        line = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = BLACK
        line.line.width = Pt(2)


def add_flow(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s, "Application Flow")
    steps = ["Splash", "Login / Signup", "OTP Verification", "Location Picker", "Home / Search", "Cart", "Checkout", "Order Tracking"]
    for i, step in enumerate(steps):
        x = 0.55 + (i % 4) * 2.85
        y = 2.25 + (i // 4) * 1.65
        box = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.15), Inches(0.62))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(232, 248, 239)
        box.line.color.rgb = GREEN
        add_text(s, step, x + 0.05, y + 0.17, 2.05, 0.25, 14, True, DARK, PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(12)
    prs.slide_height = Inches(7.5)
    add_title_slide(prs)
    simple_slide(prs, "Outline", ["Introduction", "Problem Statement", "Objectives", "Conceptual Framework", "Tools and Technologies", "System Architecture", "Modules and Implementation", "Testing and Results", "Future Scope", "References"])
    s = simple_slide(prs, "Introduction", ["quickCart is a hyperlocal quick-commerce mobile application for groceries and daily essentials.", "Users can select their state or city, browse products, add items to cart, and place Cash on Delivery orders.", "The application uses Expo SDK 54, React Native, TypeScript, Firebase, and a role-based admin panel."])
    if BANNER.exists():
        s.shapes.add_picture(str(BANNER), Inches(7.35), Inches(2.2), width=Inches(3.5))
    simple_slide(prs, "Problem Statement", ["Local grocery ordering often lacks real-time inventory visibility and simple mobile ordering.", "Small vendors need an easy way to manage products, categories, delivery cities, users, and orders.", "Customers need a fast, location-aware platform for finding essentials and tracking delivery status."])
    simple_slide(prs, "Objectives", ["Develop a cross-platform mobile app for grocery and essentials ordering.", "Implement secure authentication with email, OTP verification, and password recovery.", "Provide cart, checkout, address management, order history, and live order tracking.", "Build an admin panel for product, category, city, user, and order management."])
    simple_slide(prs, "Conceptual Framework", ["Customer side: onboarding, authentication, location selection, product discovery, cart, checkout, and order tracking.", "Admin side: dashboard, CRUD operations, user role management, and order status updates.", "Backend side: Firebase Auth, Firestore collections, Storage assets, security rules, and OTP email service."])
    simple_slide(prs, "Tools and Technologies", ["Frontend: Expo SDK 54, React Native 0.81.5, Expo Router 6, TypeScript 5.9.", "Backend: Firebase Authentication, Cloud Firestore, Firebase Storage, Firebase Functions.", "Utilities: AsyncStorage, expo-image, expo-notifications, expo-location, nodemailer OTP server.", "Build and testing: EAS Build, ESLint, TypeScript, Node test runner."])
    add_architecture(prs)
    add_flow(prs)
    simple_slide(prs, "Database / Collections", ["users: customer profiles, admin roles, push tokens.", "products: product title, price, image, category, state, city, rating, availability.", "categories: category names, icons, colors, order, visibility flags.", "orders: items, totals, payment method, delivery address, status, timestamps.", "states: delivery states, cities, nested banners/categories/products."])
    simple_slide(prs, "Main Modules", ["Authentication Module: login, signup, OTP verification, forgot password.", "Shopping Module: home feed, search, product detail, cart, checkout.", "Order Module: order placement, status workflow, order history and tracking.", "Admin Module: dashboard, products, categories, cities, users, and orders.", "Notification Module: local/user notifications and admin order alerts."])
    simple_slide(prs, "Testing and Validation", ["Unit tests validate the order status state machine.", "Firestore smoke-test scripts verify database connectivity and expected collections.", "Form validators are separated in utility modules for reuse across auth and address flows.", "Firestore security rules and role checks restrict admin functionality."])
    simple_slide(prs, "Expected Outcomes", ["A functional quick-commerce app with real shopping and checkout workflows.", "Location-aware product browsing for Indian states and cities.", "Real-time inventory and order administration for operators.", "Improved customer convenience through quick ordering and order-status visibility."])
    simple_slide(prs, "Future Scope", ["Online payment gateway integration such as UPI, cards, and wallets.", "Delivery partner app with assignment, route tracking, and proof of delivery.", "Recommendation engine based on purchase history and location trends.", "Coupons, subscriptions, inventory alerts, analytics dashboard, and multi-vendor support."])
    simple_slide(prs, "References", ["Expo SDK 54 Documentation: https://docs.expo.dev/versions/v54.0.0/", "React Native Documentation: https://reactnative.dev/", "Firebase Documentation: https://firebase.google.com/docs", "Expo Router Documentation: https://docs.expo.dev/router/introduction/", "Project source code and README: quickCart repository"])
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(s)
    add_text(s, "Thank You", 4.25, 2.65, 3.5, 0.7, 36, True, NAVY, PP_ALIGN.CENTER)
    add_text(s, "Ayush (A45304825002) - BCA", 3.7, 3.55, 4.6, 0.4, 20, False, BLACK, PP_ALIGN.CENTER)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(5.2), Inches(4.25), width=Inches(1.6))
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
