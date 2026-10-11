from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Intership persentation.pptx"
APP_LOGO = ROOT / "assets" / "images" / "quickcart-logo.png"

BLUE = RGBColor(9, 52, 98)
GOLD = RGBColor(255, 204, 0)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(246, 248, 250)
GREEN = RGBColor(34, 197, 94)


def text(slide, value, x, y, w, h, size=20, bold=False, color=BLACK, align=PP_ALIGN.LEFT, font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.space_after = Pt(9)
        p.bullet = True
    return box


def header(slide, title):
    text(slide, "AMITY\nUNIVERSITY", 0.95, 0.18, 1.55, 0.72, 23, True, BLUE, PP_ALIGN.CENTER)
    text(slide, "PATNA", 1.24, 0.82, 0.9, 0.18, 9, False, BLUE, PP_ALIGN.CENTER)
    shield = slide.shapes.add_shape(1, Inches(0.52), Inches(0.25), Inches(0.36), Inches(0.52))
    shield.fill.solid()
    shield.fill.fore_color.rgb = GOLD
    shield.line.color.rgb = BLUE
    text(slide, "A", 0.57, 0.36, 0.25, 0.18, 12, True, BLUE, PP_ALIGN.CENTER)
    slide.shapes.add_shape(1, Inches(0.72), Inches(1.05), Inches(5.2), Inches(0.06)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = GOLD
    slide.shapes.add_shape(1, Inches(5.92), Inches(1.05), Inches(5.35), Inches(0.06)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLACK
    slide.shapes.add_shape(1, Inches(0), Inches(7.42), Inches(12), Inches(0.08)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLUE
    text(slide, title, 3.35, 1.35, 5.3, 0.45, 26, True, BLACK, PP_ALIGN.CENTER)


def title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Internship Presentation")
    text(s, "React Native Developer Intern\n(AI-Assisted Development)", 3.1, 2.1, 5.8, 0.75, 24, True, BLUE, PP_ALIGN.CENTER)
    text(s, "ToonsCode Software Training Institute", 3.25, 3.0, 5.5, 0.35, 19, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Presented By", 1.1, 4.0, 3.7, 0.35, 18, True, BLACK, PP_ALIGN.CENTER)
    text(s, "Ayush (A45304825002) - BCA", 1.0, 4.42, 3.9, 0.35, 18, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Under the Supervision of", 7.05, 4.0, 3.6, 0.35, 18, True, BLACK, PP_ALIGN.CENTER)
    text(s, "Dr. Sushant Kumar Dubey", 7.05, 4.42, 3.6, 0.35, 18, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Amity Institute Of Information Technology, Patna\n10 August, 2026", 3.2, 5.85, 5.6, 0.55, 18, False, BLACK, PP_ALIGN.CENTER)


def simple(prs, title, items):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, title)
    bullets(s, items, 1.15, 2.2, 9.7, 4.5)


def certificate_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Internship Certificate")
    card = s.shapes.add_shape(1, Inches(1.0), Inches(1.95), Inches(10.0), Inches(4.75))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BLUE
    card.line.width = Pt(2)
    text(s, "ToonsCode", 4.55, 2.18, 2.9, 0.35, 24, True, BLACK, PP_ALIGN.CENTER)
    text(s, "(a unit of Chapterfeed Learning Space PVT. LTD.)", 4.15, 2.52, 3.7, 0.22, 8, False, BLACK, PP_ALIGN.CENTER)
    text(s, "CERTIFICATE", 3.38, 2.8, 5.25, 0.55, 36, True, BLUE, PP_ALIGN.CENTER)
    bar = s.shapes.add_shape(1, Inches(3.6), Inches(3.34), Inches(4.8), Inches(0.26))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.color.rgb = BLUE
    text(s, "OF  INTERNSHIP", 3.65, 3.37, 4.7, 0.18, 12, True, WHITE, PP_ALIGN.CENTER)
    text(s, "This certificate is proudly presented to", 3.35, 3.75, 5.3, 0.26, 15, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Mr. AYUSH", 3.55, 4.12, 4.9, 0.45, 28, True, BLACK, PP_ALIGN.CENTER)
    text(s, "In recognition of successful completion of the 6-week internship program as a", 2.7, 4.66, 6.6, 0.22, 10, False, BLACK, PP_ALIGN.CENTER)
    text(s, "React Native Developer Intern (AI-Assisted Development)", 2.35, 4.95, 7.3, 0.28, 16, True, BLACK, PP_ALIGN.CENTER)
    text(s, "Conducted from 20th June 2026 to 2nd August 2026", 3.35, 5.28, 5.3, 0.22, 10, True, BLACK, PP_ALIGN.CENTER)
    text(s, "Gained practical experience in React Native, JavaScript/TypeScript, React,\nFirebase integration, REST APIs, and AI-assisted coding practices.", 2.1, 5.58, 7.8, 0.38, 10, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Bikalp Anand\nDirector", 1.75, 6.05, 1.7, 0.45, 10, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Abhishek Kumar Singh\nInstructor", 8.5, 6.05, 1.9, 0.45, 10, False, BLACK, PP_ALIGN.CENTER)


def offer_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Offer Letter")
    page = s.shapes.add_shape(1, Inches(2.35), Inches(1.78), Inches(7.3), Inches(5.55))
    page.fill.solid()
    page.fill.fore_color.rgb = WHITE
    page.line.color.rgb = RGBColor(220, 225, 220)
    page.line.width = Pt(1)
    text(s, "ToonsCode", 4.4, 1.98, 3.2, 0.32, 23, True, BLACK, PP_ALIGN.CENTER)
    text(s, "(a unit of Chapterfeed Learning Space Pvt. Ltd.)\nSoftware Training Institute", 4.0, 2.28, 4.0, 0.32, 8, True, BLACK, PP_ALIGN.CENTER)
    text(s, "Ref. No.: TC/DEV/2026/101/IS\nDate: 15/06/2026", 7.65, 2.58, 1.55, 0.34, 7, False, BLACK, PP_ALIGN.LEFT)
    text(s, "Appointment for Intern - React Native Developer Intern (AI-Assisted Development)", 2.85, 3.08, 6.2, 0.18, 8, True, BLACK)
    body = (
        "Dear Mr. Ayush,\n\n"
        "Congratulations on your selection for the Certificate in React Native with AI-Assisted Coding Internship Program at ToonsCode.\n\n"
        "The 6-week internship commenced on 20th June 2026. During the internship, the role focused on cross-platform mobile application development using React Native, JavaScript/TypeScript, React, reusable components, REST APIs, debugging, and AI-assisted coding practices.\n\n"
        "The internship provided exposure to contemporary mobile development technologies, software engineering principles, AI-assisted workflows, deployment practices, and collaborative project environments.\n\n"
        "Once again, congratulations on your selection. We look forward to your valuable contributions as a React Native Developer Intern (AI-Assisted Development) at ToonsCode.\n\n"
        "Acceptance of Our Offer Letter\nPlease confirm your acceptance of this Contract of Internship by signing and returning the duplicate copy."
    )
    text(s, body, 2.85, 3.34, 6.2, 2.4, 7, False, BLACK, PP_ALIGN.LEFT)
    text(s, "Bikalp\nAnand", 2.95, 5.86, 1.0, 0.35, 12, False, BLACK, PP_ALIGN.CENTER, "Segoe Script")
    left_line = s.shapes.add_shape(1, Inches(2.8), Inches(6.18), Inches(1.15), Inches(0.02))
    left_line.fill.solid()
    left_line.fill.fore_color.rgb = BLACK
    left_line.line.color.rgb = BLACK
    text(s, "Warm Regards,\nBikalp Anand\nFounder\nToonsCode", 2.85, 6.22, 1.45, 0.42, 6, False, BLACK)
    text(s, "Ayush", 6.55, 5.72, 2.2, 0.55, 28, False, BLACK, PP_ALIGN.CENTER, "Segoe Script")
    line = s.shapes.add_shape(1, Inches(6.75), Inches(6.18), Inches(1.55), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.color.rgb = BLACK
    text(s, "Mr. Ayush\nReact Native Developer Intern", 6.65, 6.22, 1.8, 0.28, 6, False, BLACK, PP_ALIGN.CENTER)
    footer = s.shapes.add_shape(1, Inches(2.35), Inches(6.78), Inches(7.3), Inches(0.25))
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(25, 32, 32)
    footer.line.color.rgb = RGBColor(25, 32, 32)
    text(s, "88 77 10 1234     connect@toonscode.com     www.toonscode.com", 2.45, 6.85, 7.1, 0.1, 6, True, WHITE, PP_ALIGN.CENTER)
    footer2 = s.shapes.add_shape(1, Inches(2.35), Inches(7.06), Inches(7.3), Inches(0.22))
    footer2.fill.solid()
    footer2.fill.fore_color.rgb = RGBColor(25, 32, 32)
    footer2.line.color.rgb = RGBColor(25, 32, 32)
    text(s, "Toonscode, Mangalam Road, Judges Colony, RPS, Danapur, Patna - 801503", 2.55, 7.11, 6.9, 0.09, 6, True, WHITE, PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(12)
    prs.slide_height = Inches(7.5)

    title_slide(prs)
    simple(prs, "Outline", [
        "Internship introduction",
        "Organization profile",
        "Offer letter and joining details",
        "Certificate details",
        "Work completed during internship",
        "Learning outcomes",
        "Project contribution",
        "Conclusion",
    ])
    simple(prs, "Internship Introduction", [
        "Internship role: React Native Developer Intern (AI-Assisted Development).",
        "Duration: 20th June 2026 to 2nd August 2026.",
        "Organization: ToonsCode Software Training Institute.",
        "Focus area: cross-platform mobile application development using React Native and AI-assisted coding practices.",
    ])
    simple(prs, "Organization Profile", [
        "ToonsCode is a software training institute under Chapterfeed Learning Space Private Limited.",
        "The internship emphasized practical mobile app development, reusable UI components, REST APIs, Firebase integration, debugging, and deployment workflows.",
        "Mentorship was provided by experienced professionals during the training period.",
    ])
    offer_slide(prs)
    certificate_slide(prs)
    simple(prs, "Work Completed", [
        "Designed and improved mobile application screens using React Native and TypeScript.",
        "Implemented and tested reusable components for authentication, product browsing, cart, checkout, and order tracking.",
        "Integrated Firebase services for authentication, database storage, and app data management.",
        "Used AI-assisted development tools to improve productivity and debugging efficiency.",
    ])
    simple(prs, "Learning Outcomes", [
        "Improved understanding of Expo, React Native, JavaScript, TypeScript, and React concepts.",
        "Learned Firebase Authentication, Firestore data modeling, Storage, and security-rule basics.",
        "Gained exposure to REST APIs, mobile UI design, Git/version control, testing, and deployment practices.",
        "Developed better problem-solving and project documentation skills.",
    ])
    simple(prs, "Project Contribution", [
        "Applied internship learning to the quickCart mobile application.",
        "Built customer-facing workflows such as login, location selection, product listing, cart, checkout, and orders.",
        "Created admin workflows for products, categories, cities, users, and order-status management.",
        "Focused on practical quick-commerce use cases for groceries and daily essentials.",
    ])
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Conclusion")
    text(s, "The internship helped bridge academic knowledge and practical mobile app development.", 1.4, 2.55, 9.2, 0.35, 21, False, BLACK, PP_ALIGN.CENTER)
    text(s, "Thank You", 4.15, 3.45, 3.7, 0.55, 34, True, BLUE, PP_ALIGN.CENTER)
    text(s, "Ayush (A45304825002) - BCA", 3.65, 4.15, 4.7, 0.35, 19, False, BLACK, PP_ALIGN.CENTER)
    if APP_LOGO.exists():
        s.shapes.add_picture(str(APP_LOGO), Inches(5.25), Inches(4.75), width=Inches(1.5))

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
