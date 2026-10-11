from pptx import Presentation

prs = Presentation('Intership persentation.pptx')
print(f'Total slides: {len(prs.slides)}')
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()[:60]
            if t:
                texts.append(t)
    print(f'Slide {i+1}: {" | ".join(texts[:4])}')
