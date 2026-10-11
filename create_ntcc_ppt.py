from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = "ntcc ppt.pptx"

BG = RGBColor(247, 249, 252)
NAVY = RGBColor(24, 43, 77)
TEAL = RGBColor(0, 145, 124)
GREEN = RGBColor(42, 156, 88)
ORANGE = RGBColor(238, 132, 58)
RED = RGBColor(196, 64, 64)
MUTED = RGBColor(92, 105, 122)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(232, 239, 246)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_header(slide, title, subtitle=None):
    add_text(slide, "AMITY UNIVERSITY PATNA", 0.55, 0.24, 3.2, 0.28, 9, MUTED, True)
    add_text(slide, "quickCart NTCC Progress Report", 8.0, 0.24, 3.4, 0.28, 9, MUTED, False, PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.66, 7.8, 0.46, 24, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.57, 1.12, 8.8, 0.3, 10.5, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.48), Inches(10.9), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()


def add_bullet_box(slide, items, x, y, w, h, title=None, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LIGHT
    shape.adjustments[0] = 0.08
    top = y + 0.2
    if title:
        add_text(slide, title, x + 0.25, top, w - 0.5, 0.3, 13, accent, True)
        top += 0.42
    box = slide.shapes.add_textbox(Inches(x + 0.28), Inches(top), Inches(w - 0.52), Inches(h - (top - y) - 0.18))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(6)
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY
    return shape


def add_pill(slide, text, x, y, w, color):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.35))
    pill.fill.solid()
    pill.fill.fore_color.rgb = color
    pill.line.fill.background()
    pill.adjustments[0] = 0.45
    add_text(slide, text, x, y + 0.06, w, 0.2, 8.5, WHITE, True, PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = Inches(12)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 1. Title
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_text(slide, "NTCC Progress Report", 0.65, 0.95, 6.6, 0.6, 30, NAVY, True)
add_text(slide, "quickCart: Hyperlocal Grocery & Essentials Delivery App", 0.68, 1.58, 7.4, 0.5, 17, TEAL, True)
add_text(slide, "Weekly Progress Report, Conceptual Framework, Objectives & Methodology", 0.7, 2.15, 7.0, 0.45, 13, MUTED)
add_pill(slide, "Expo SDK 54", 0.72, 2.85, 1.25, NAVY)
add_pill(slide, "React Native", 2.12, 2.85, 1.3, TEAL)
add_pill(slide, "Firebase", 3.56, 2.85, 1.05, ORANGE)
add_pill(slide, "TypeScript", 4.75, 2.85, 1.1, GREEN)
add_bullet_box(slide, [
    "Student: Ayush",
    "Project type: Mobile application development",
    "Domain: Quick commerce / e-commerce",
    "Platform: Android and iOS"
], 0.7, 3.55, 4.8, 2.1, "Project Information", TEAL)
slide.shapes.add_picture("assets/images/readme_banner.jpg", Inches(6.15), Inches(1.05), width=Inches(5.1))

# 2. Objectives
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Objectives", "What the NTCC project aims to achieve")
add_bullet_box(slide, [
    "Design a user-friendly quick-commerce mobile app for groceries and daily essentials.",
    "Implement authentication, OTP verification, location selection, product browsing and cart flow.",
    "Enable order placement with real-time status tracking through Firebase Firestore.",
    "Build an admin panel for product, category, city, order and user management.",
    "Apply maintainable Expo Router, TypeScript and Firebase architecture."
], 0.8, 1.85, 10.4, 4.7, "Primary Objectives", GREEN)

# 3. Conceptual Framework
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Conceptual Framework", "How users, app modules and backend services interact")
steps = [
    ("User", 0.9, 2.5, TEAL),
    ("Mobile App\nExpo + RN", 3.0, 2.5, NAVY),
    ("Firebase\nAuth/Firestore", 5.25, 2.5, ORANGE),
    ("Admin Panel", 7.75, 2.5, GREEN),
    ("Order\nFulfilment", 9.75, 2.5, RED),
]
for text, x, y, color in steps:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.55), Inches(0.85))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    box.adjustments[0] = 0.15
    add_text(slide, text, x + 0.05, y + 0.18, 1.45, 0.45, 11, WHITE, True, PP_ALIGN.CENTER)
for x in [2.55, 4.75, 7.2, 9.25]:
    add_text(slide, "->", x, 2.72, 0.35, 0.25, 16, MUTED, True, PP_ALIGN.CENTER)
add_bullet_box(slide, [
    "Input: user credentials, selected location, product demand and cart data.",
    "Processing: validation, product filtering, cart calculation and order status updates.",
    "Storage: users, products, categories, states/cities, banners and orders in Firebase.",
    "Output: personalized product feed, confirmed order, live status and admin reports."
], 1.25, 4.1, 9.5, 2.05, "Framework Summary", TEAL)

# 4. Methodology
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Methodology", "Development approach followed during the project")
add_bullet_box(slide, [
    "Requirement analysis: studied quick-commerce flows, user roles and Firebase data needs.",
    "UI/UX planning: mapped splash, auth, location, home, search, cart, checkout and profile screens.",
    "Implementation: built Expo Router screens, reusable components and TypeScript utilities.",
    "Backend integration: configured Firebase Auth, Firestore, Storage and custom email OTP server.",
    "Testing and refinement: validated order status logic, Firestore rules and core user journeys."
], 0.8, 1.85, 10.4, 4.7, "Iterative Development Method", ORANGE)

# 5. Weekly Progress
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Weekly Progress Report", "Summary of work completed across project phases")
rows = [
    ("Week 1", "Requirement study, scope definition, project planning and technology selection."),
    ("Week 2", "Expo SDK 54 setup, routing structure, Firebase configuration and UI flow planning."),
    ("Week 3", "Authentication, OTP verification, permission flow and state/location picker."),
    ("Week 4", "Home feed, category rail, search, product detail, cart and checkout modules."),
    ("Week 5", "Order placement, order status tracking and admin dashboard/product management."),
    ("Week 6", "Admin order/user management, Firestore rules, testing, documentation and final review.")
]
y = 1.85
for i, (week, work) in enumerate(rows):
    color = TEAL if i % 2 == 0 else NAVY
    add_pill(slide, week, 0.85, y, 1.05, color)
    add_text(slide, work, 2.1, y + 0.04, 8.8, 0.28, 10.5, NAVY)
    y += 0.55

# 6. Current Status
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Current Status", "Present implementation progress as of 10 August 2026")
add_bullet_box(slide, [
    "Functional prototype stage with customer and admin workflows implemented.",
    "Customer screens include auth, OTP, location picker, home, search, product detail, cart, checkout, orders and profile.",
    "Admin screens include dashboard, products, categories, cities, orders, users and CRUD management.",
    "Firebase configuration, Firestore rules, indexes, storage rules and service modules are present.",
    "Project documentation and presentation/reporting assets have been prepared."
], 0.75, 1.85, 5.1, 4.65, "Completed Progress", TEAL)
add_bullet_box(slide, [
    "Expo SDK 54 with React Native 0.81.5 and TypeScript 5.9.",
    "Expo Router 6 used for file-based customer and admin navigation.",
    "Firebase Auth, Firestore and Storage used for backend workflows.",
    "Reusable utilities handle cart state, selected state, validation, currency and order status.",
    "EAS build profiles are configured for development, preview and production."
], 6.15, 1.85, 5.1, 4.65, "Technical Progress", GREEN)

# 7. Modules
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Modules & System Design", "Major functional blocks of quickCart")
add_bullet_box(slide, [
    "Customer app: onboarding, auth, browsing, cart, checkout, orders and profile.",
    "Admin app: dashboard, inventory CRUD, city setup, order updates and user roles.",
    "Data layer: Firestore services, typed utilities and lightweight local state stores.",
    "Security: role-gated admin routes and Firestore security rules."
], 0.75, 1.85, 5.1, 4.65, "Functional Modules", TEAL)
add_bullet_box(slide, [
    "users: customer profile and role.",
    "products: product catalog by state and category.",
    "categories: category metadata and display ordering.",
    "orders: cart items, customer details and delivery status.",
    "states: serviceable states and cities."
], 6.15, 1.85, 5.1, 4.65, "Database Structure", GREEN)

# 8. Verification & Next Steps
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Verification & Next Steps", "Testing result, known issue and planned improvements")
add_bullet_box(slide, [
    "Lint command executed successfully with 0 errors and 33 warnings.",
    "Warnings mainly relate to unused imports/variables and React hook dependency checks.",
    "Unit test command executed but 2 order-status tests failed.",
    "Failure reason: test expectations use an older status flow than the current implementation."
], 0.75, 1.85, 5.1, 4.65, "Verification Status", ORANGE)
add_bullet_box(slide, [
    "Align order-status tests with the final business workflow.",
    "Clean lint warnings and rerun verification commands.",
    "Run customer flow testing from signup to checkout.",
    "Run admin flow testing for product management and order status updates.",
    "Review generated presentation/report files before final submission."
], 6.15, 1.85, 5.1, 4.65, "Next Steps", RED)

# 9. Expected Outcomes
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Expected Outcomes", "Deliverables and benefits of the project")
add_bullet_box(slide, [
    "A working quick-commerce mobile application for grocery and essentials ordering.",
    "Admin workflow to keep products, cities, users and orders manageable.",
    "Firebase-backed real-time order tracking and scalable product data storage.",
    "Cleaner understanding of mobile app architecture, backend integration and deployment workflow.",
    "NTCC documentation and presentation artifacts for academic evaluation."
], 0.8, 1.85, 10.4, 4.7, "Deliverables", RED)

# 10. Conclusion
slide = prs.slides.add_slide(blank)
set_bg(slide)
add_header(slide, "Conclusion", "Current status and next steps")
add_bullet_box(slide, [
    "Core customer journey and admin management workflows are defined and implemented.",
    "The project demonstrates practical use of Expo, React Native, Firebase and TypeScript.",
    "The main stabilization task is to align order-status tests with the implemented workflow.",
    "Next improvements include payment integration, delivery partner workflow, analytics and production hardening."
], 1.2, 2.05, 9.6, 3.0, "Project Summary", TEAL)
add_text(slide, "Thank You", 0.0, 5.8, 12.0, 0.5, 24, NAVY, True, PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Created {OUT} with {len(prs.slides)} slides")
